import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_project_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Executive Theme Color Palette
    BG_DARK = RGBColor(10, 17, 40)       # #0A1128 Deep Slate Navy
    CARD_BG = RGBColor(19, 31, 55)       # #131F37 Rich Dark Container
    CARD_BORDER = RGBColor(42, 59, 92)   # #2A3B5C Muted Border
    HEADER_FILL = RGBColor(27, 42, 74)   # #1B2A4A Table / Badge Fill
    
    C_BLUE = RGBColor(0, 180, 216)       # #00B4D8 Electric Sky
    C_EMERALD = RGBColor(16, 185, 129)   # #10B981 Vivid Emerald
    C_CORAL = RGBColor(255, 90, 95)      # #FF5A5F Vibrant Coral
    C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Golden Amber
    C_PURPLE = RGBColor(168, 85, 247)    # #A855F7 Purple Accent
    C_WHITE = RGBColor(255, 255, 255)
    C_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Cool Grey

    FONT = "Segoe UI"

    def apply_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    def add_header(slide, title, section_name=""):
        apply_slide_bg(slide)

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"SKIP-Q  |  TECHNICAL & FUNCTIONAL PROJECT DEEP DIVE   •   {section_name.upper()}"
        p0.font.name = FONT
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_before = Pt(4)

    # ==========================================
    # SLIDE 1: COVER SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s1)

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = C_BLUE
    c1.line.width = Pt(2)

    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.6)
    tf.margin_right = Inches(0.6)

    p = tf.paragraphs[0]
    p.text = "COMPLETE SYSTEM ARCHITECTURE & WORKING DEMO"
    p.font.name = FONT
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Skip-Q"
    p.font.name = FONT
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_before = Pt(10)

    p = tf.add_paragraph()
    p.text = "Smart OPD Queue Management & Live In-Room Consultation Radar"
    p.font.name = FONT
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "A complete, end-to-end full stack system eliminating physical waiting room lines through real-time consultation tracking."
    p.font.name = FONT
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = C_MUTED
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Live Portals:\n• Patient App: https://skipq-user.vercel.app\n• Hospital Portal: https://skipq-hospital.vercel.app\n• Admin Hub: https://skipq-admin.vercel.app"
    p.font.name = FONT
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p.space_before = Pt(24)

    # ==========================================
    # SLIDE 2: THE CORE IDEA & GENESIS
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Core Idea: Why Skip-Q was Built", "01. Concept & Vision")

    idea_cards = [
        ("The Real Problem", "Patients in Tier-2/3 towns travel 30+ km from rural mandals, only to spend 2 to 4 hours waiting blindly on crowded hospital benches without knowing when the doctor will arrive or which token is being treated.", C_CORAL),
        ("The Core Concept", "Skip-Q transforms the hospital OPD waiting experience into a 'Live Radar' on the patient's smartphone. Patients can track the doctor's ongoing consultation token in real-time and arrive right on time.", C_BLUE),
        ("The Ultimate Goal", "Zero physical waiting room chaos, 40% faster clinic patient throughput, zero cross-infection risks in crowded lobbies, and complete digital transparency for doctors and patients alike.", C_EMERALD)
    ]

    for i, (title, desc, col) in enumerate(idea_cards):
        x = Inches(0.8 + i * 4.04)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.name = FONT
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(14)

    # ==========================================
    # SLIDE 3: COMPLETE END-TO-END FLOW OVERVIEW
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "End-to-End System Workflow: Initial to End", "02. Complete Workflow")

    steps = [
        ("Step 1: Discovery", "Patient opens app. Location radar auto-detects GPS / Mahabubabad town boundary and displays local hospitals with live token boards."),
        ("Step 2: Digital Booking", "Patient selects doctor, enters contact details, verifies 6-digit OTP, and books a real-time digital OPD token pass in seconds."),
        ("Step 3: Live Radar", "Patient tracks doctor's room status live from home or tea stall (e.g., 'Now Serving #12 | Your Token #15 | Approx 15 Mins')."),
        ("Step 4: Reception Desk", "Receptionist clicks 'Call Next Patient' on the Hospital Portal. System automatically alerts the patient and advances the queue."),
        ("Step 5: Consultation & Finish", "Doctor completes consultation. Next patient enters smoothly with zero lobby shouting or door knocking.")
    ]

    for i, (st_title, st_desc) in enumerate(steps):
        y = Inches(1.65 + i * 1.05)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.92))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = C_BLUE if i % 2 == 0 else C_EMERALD
        c.line.width = Pt(1.2)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = f"{st_title}:  {st_desc}"
        p.font.name = FONT
        p.font.size = Pt(12)
        p.font.color.rgb = C_WHITE

    # ==========================================
    # SLIDE 4: THE 3-PORTAL ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "The 3 Dedicated Portals & Their Roles", "03. Architecture")

    portals = [
        ("📱 Patient Portal", "skipq-user.vercel.app", C_BLUE, [
            "Hyperlocal Hospital Discovery",
            "Doctor profiles & OPD timings",
            "10-Sec OTP Token Generation",
            "Live Countdown Queue Radar",
            "SMS & Email turn notifications"
        ]),
        ("🏢 Hospital Portal", "skipq-hospital.vercel.app", C_EMERALD, [
            "1-Click 'Call Next Patient' desk",
            "Live queue caller controller",
            "Emergency token skip & recall",
            "Doctor OPD schedule manager",
            "Daily patient traffic analytics"
        ]),
        ("👑 Super Admin Hub", "skipq-admin.vercel.app", C_PURPLE, [
            "District hospital onboarding",
            "Doctor credential & bio approval",
            "Multi-hospital live queue audit",
            "Supabase database governance",
            "Security & platform protection"
        ])
    ]

    for i, (p_title, p_url, p_col, p_features) in enumerate(portals):
        x = Inches(0.8 + i * 4.04)
        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = p_col
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.35)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.name = FONT
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        p = tf.add_paragraph()
        p.text = p_url
        p.font.name = FONT
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = p_col
        p.space_before = Pt(2)

        for feat in p_features:
            p = tf.add_paragraph()
            p.text = f"✔ {feat}"
            p.font.name = FONT
            p.font.size = Pt(12)
            p.font.color.rgb = C_MUTED
            p.space_before = Pt(8)

    # ==========================================
    # SLIDE 5: PATIENT JOURNEY (DETAILED)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Deep Dive: The Patient Experience", "04. Patient App Functionality")

    p_flow = [
        ("1. Geolocation & Clinic Discovery", "Patient opens app. Integrated Photon/OpenStreetMap engine automatically searches clinics in Mahabubabad. Filters doctors by specialization and real-time availability."),
        ("2. 4-Step Secure Authentication", "Step 1: Contact details (strict 10-digit phone & domain-verified email) ➔ Step 2: 6-digit Email OTP with 5-attempt guard ➔ Step 3: Password setup ➔ Step 4: Mandatory town selection."),
        ("3. Instant Token Issuance", "System allocates a unique digital token ID (e.g. #18) linked to the doctor's queue in Supabase PostgreSQL database with timestamp and token status."),
        ("4. Real-Time Consultation Radar", "Patient views live countdown ('3 Patients Ahead of You | Avg 5 mins/patient'). Patient receives SMS alert when 2 patients away, arriving at clinic exactly on time.")
    ]

    for i, (title, desc) in enumerate(p_flow):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.96)
        y = Inches(1.65 + row * 2.65)

        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.77), Inches(2.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = C_BLUE
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.3)
        tf.margin_right = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_BLUE

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 6: HOSPITAL DESK & DOCTOR MANAGEMENT
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Deep Dive: Hospital Desk & Doctor Caller Desk", "05. Hospital Functionality")

    h_flow = [
        ("1. Reception Desk Queue Caller", "Ultra-simplified 1-click interface. Receptionist clicks 'Call Next' to advance queue. Live token screen outside doctor's room syncs automatically."),
        ("2. Queue Control Actions", "• Recall Patient (if patient stepped out)\n• Skip / Push Back (for late arrivals)\n• Mark Emergency Priority (immediate bypass)\n• End OPD Session."),
        ("3. Doctor Schedule & Bio Control", "Hospital administrator adds doctor specializations, consulting hours, OPD room numbers, and consultation fees with instant live updates."),
        ("4. Zero Hardware Capital Expense", "Runs on any existing receptionist desktop, tablet, or smartphone browser. Zero kiosks or proprietary display hardware required.")
    ]

    for i, (title, desc) in enumerate(h_flow):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.96)
        y = Inches(1.65 + row * 2.65)

        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.77), Inches(2.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = C_EMERALD
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.3)
        tf.margin_right = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(12.5)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 7: SUPER ADMIN & GOVERNANCE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Deep Dive: Super Admin & District Governance", "06. Admin Functionality")

    admin_cards = [
        ("Hospital Onboarding Engine", "Super Admin registers new hospitals, verifies clinic addresses, sets geofences, and provisions initial administrator credentials in under 5 minutes.", C_BLUE),
        ("Doctor Credential Verification", "Admin approves doctor registrations, ensures medical licenses are active, and assigns specialized OPD departments across district clinics.", C_PURPLE),
        ("District-Wide Live Queue Radar", "Master supervision dashboard monitoring real-time active patient traffic across all registered hospitals in Mahabubabad district simultaneously.", C_EMERALD)
    ]

    for i, (title, desc, col) in enumerate(admin_cards):
        x = Inches(0.8 + i * 4.04)
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = f"★ {title}"
        p.font.name = FONT
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(14)

    # ==========================================
    # SLIDE 8: TECH STACK & SYNCHRONIZATION
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Technical Architecture & Real-Time Sync", "07. Technology Stack")

    tech_table_shape = s8.shapes.add_table(5, 3, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.15))
    tt = tech_table_shape.table
    tt.columns[0].width = Inches(3.2)
    tt.columns[1].width = Inches(3.8)
    tt.columns[2].width = Inches(4.733)

    t_headers = ["Layer / Component", "Technology Used", "Role in Skip-Q"]
    for i, h in enumerate(t_headers):
        cell = tt.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_FILL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_BLUE
        p.font.name = FONT

    t_rows = [
        ("Frontend & UI Framework", "Next.js 14 (App Router) + TailwindCSS", "Ultra-fast server-rendered UI, responsive mobile design, dynamic dark/light theme."),
        ("Database & ORM", "Supabase PostgreSQL + Prisma ORM", "Persistent relational storage for users, tokens, doctor queues, and OTP state."),
        ("Real-Time Synchronization", "WebSockets + Next.js Serverless Routes", "Sub-second token updates between Hospital Caller and Patient Radar."),
        ("Security & Anti-Inspection", "Custom SecurityGuard & SHA-256 Hashing", "Blocks F12, DevTools hotkeys, right-click inspection, and secures auth tokens.")
    ]

    for r_idx, (layer, tech, role) in enumerate(t_rows, start=1):
        for c_idx, val in enumerate([layer, tech, role]):
            cell = tt.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = C_WHITE if c_idx == 0 else (C_BLUE if c_idx == 1 else C_MUTED)
            p.font.bold = (c_idx == 0 or c_idx == 1)
            p.font.name = FONT

    # ==========================================
    # SLIDE 9: KEY INNOVATIONS & SUMMARY
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Summary of Key Innovations & Real-World Impact", "08. Key Innovations")

    summary_cards = [
        ("Zero Physical Waiting", "Transforms crowded waiting rooms into a digital mobile queue with live countdown timers and SMS alerts.", C_BLUE),
        ("Zero CapEx Onboarding", "Runs directly on existing clinic hardware in under 10 minutes without requiring ₹1.5L kiosk machines.", C_EMERALD),
        ("Hyperlocal Focus", "Engineered specifically for Tier-2/3 district towns where Practo and corporate platforms do not operate.", C_AMBER),
        ("100% Live & Functional", "Fully deployed on Vercel, Supabase PostgreSQL, and open source for demonstration and immediate rollout.", C_PURPLE)
    ]

    for i, (stitle, sdesc, scol) in enumerate(summary_cards):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.96)
        y = Inches(1.65 + row * 2.65)

        c = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.77), Inches(2.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = scol
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.3)
        tf.margin_right = Inches(0.35)

        p = tf.paragraphs[0]
        p.text = f"✔ {stitle}"
        p.font.name = FONT
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = scol

        p = tf.add_paragraph()
        p.text = sdesc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 10: CLOSING & LIVE DEMO
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s10)

    c10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    c10.fill.solid()
    c10.fill.fore_color.rgb = CARD_BG
    c10.line.color.rgb = C_EMERALD
    c10.line.width = Pt(2)

    tf = c10.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.6)
    tf.margin_right = Inches(0.6)

    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = FONT
    p.font.size = Pt(52)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = '"Eliminating Hospital Waiting Room Chaos through Technology"'
    p.font.name = FONT
    p.font.size = Pt(18)
    p.font.italic = True
    p.font.color.rgb = C_WHITE
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Live Interactive Prototype Demonstration:\n• 📱 Patient Web App: https://skipq-user.vercel.app\n• 🏢 Hospital Reception Desk: https://skipq-hospital.vercel.app\n• 👑 Super Admin Governance Hub: https://skipq-admin.vercel.app\n• 💻 Full Open Source Code: https://github.com/sameir-dev/Skip-Q.git"
    p.font.name = FONT
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_before = Pt(26)

    # Output file paths
    out_path = "C:/OnlineConsultation/Skip_Q_Project_Working_and_Idea.pptx"
    prs.save(out_path)
    print(f"Complete Project Idea & Functionality Presentation created at: {out_path}")

if __name__ == "__main__":
    build_project_deck()

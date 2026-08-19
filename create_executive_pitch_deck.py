import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Executive Theme Colors
    BG_DARK = RGBColor(10, 17, 40)       # #0A1128 Deep Slate Navy
    CARD_BG = RGBColor(19, 31, 55)       # #131F37 Rich Dark Container
    CARD_BORDER = RGBColor(42, 59, 92)   # #2A3B5C Muted Border
    HEADER_FILL = RGBColor(27, 42, 74)   # #1B2A4A Table / Badge Fill
    
    C_BLUE = RGBColor(0, 180, 216)       # #00B4D8 Electric Sky
    C_EMERALD = RGBColor(16, 185, 129)   # #10B981 Vivid Emerald
    C_CORAL = RGBColor(255, 90, 95)      # #FF5A5F Vibrant Coral
    C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Golden Amber
    C_WHITE = RGBColor(255, 255, 255)
    C_LIGHT = RGBColor(226, 232, 240)    # #E2E8F0 Soft White
    C_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Cool Grey

    FONT = "Segoe UI"

    def apply_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    def add_top_bar(slide, title, category=""):
        apply_slide_bg(slide)

        # Header Container
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"SKIP-Q  |  JNTUH TBI IDEATHON 2026   •   {category.upper()}"
        p0.font.name = FONT
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 1: COVER SLIDE
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s1)

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = C_BLUE
    c1.line.width = Pt(2)

    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.6)
    tf.margin_right = Inches(0.6)

    p = tf.paragraphs[0]
    p.text = "JNTUH TBI IDEATHON 2026   •   STARTUP PITCH"
    p.font.name = FONT
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Skip-Q"
    p.font.name = FONT
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_before = Pt(10)

    p = tf.add_paragraph()
    p.text = "Smart OPD Queue Management & Live In-Room Consultation Radar"
    p.font.name = FONT
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Restoring Time and Dignity to Healthcare Waiting Rooms across Bharat's Tier-2 & Tier-3 Towns"
    p.font.name = FONT
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = C_MUTED
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Founder: Sameer Shaik   •   JNTUH TBI 2-Day Bootcamp (20–21 August 2026)\nLive Web Platform: https://skipq-user.vercel.app"
    p.font.name = FONT
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p.space_before = Pt(28)

    # ----------------------------------------------------
    # SLIDE 2: THE PROBLEM
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_top_bar(s2, "The Agony of Indian Hospital OPDs", "01. The Problem")

    problems = [
        ("01", "2 to 4 Hours Wasted", "Patients travel 30+ km from rural mandals and sit blindly on crowded hospital benches without knowing when their doctor will call them.", C_CORAL),
        ("02", "70% Bharat Neglected", "Metro apps like Practo focus only on high-fee scheduled appointments, completely ignoring walk-in OPD queues in district towns like Mahabubabad.", C_AMBER),
        ("03", "Cross-Infection & Chaos", "Overcrowded waiting halls increase airborne viral transmission, noise pollution, doctor burnout, and receptionist exhaustion.", C_CORAL)
    ]

    for i, (num, title, desc, col) in enumerate(problems):
        x = Inches(0.8 + i * 4.04)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = num
        p.font.name = FONT
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(14)

    # ----------------------------------------------------
    # SLIDE 3: THE SOLUTION
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_top_bar(s3, "Live In-Room Consultation Radar", "02. The Solution")

    solutions = [
        ("📡 Live In-Room Token Radar", "Patients track the exact ongoing token number live on their mobile (Now Serving: #14 | Your Token: #18) with sub-second WebSocket updates."),
        ("🎟️ 10-Second Digital Token Pass", "Book verified OPD passes in seconds using phone & email OTP, complete with live estimated arrival countdown timers."),
        ("💻 Zero Hardware Capital Cost", "Zero expensive ₹1.5L kiosk machines needed. Clinics run the entire caller desk directly from any smartphone, tablet, or laptop browser."),
        ("📍 Hyperlocal Town Radar", "Sub-100ms OpenStreetMap location detection tailored specifically for district headquarters like Mahabubabad.")
    ]

    for i, (title, desc) in enumerate(solutions):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.96)
        y = Inches(1.65 + row * 2.65)

        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.77), Inches(2.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = C_BLUE
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_BLUE

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 4: ARCHITECTURE
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_top_bar(s4, "Unified 3-Pillar Cloud Infrastructure", "03. Architecture")

    archs = [
        ("1. Patient Web App", "skipq-user.vercel.app", C_BLUE, [
            "Live clinic discovery & location radar",
            "Real-time token counter & countdown",
            "Instant OTP verification & token pass",
            "SMS/Email updates before your turn"
        ]),
        ("2. Hospital Desk Portal", "skipq-hospital.vercel.app", C_EMERALD, [
            "1-Click 'Call Next Patient' desk",
            "Doctor schedule & OPD session manager",
            "Emergency token priority override",
            "Daily patient throughput analytics"
        ]),
        ("3. Super Admin Hub", "skipq-admin.vercel.app", C_BLUE, [
            "District-wide hospital onboarding",
            "Doctor credential & bio verification",
            "Live system-wide consultation tracker",
            "Supabase PostgreSQL governance"
        ])
    ]

    for i, (title, url, col, bullets) in enumerate(archs):
        x = Inches(0.8 + i * 4.04)
        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.35)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        p = tf.add_paragraph()
        p.text = url
        p.font.name = FONT
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = col
        p.space_before = Pt(2)

        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"• {b}"
            p.font.name = FONT
            p.font.size = Pt(12)
            p.font.color.rgb = C_MUTED
            p.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 5: BUSINESS MODEL CANVAS
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_top_bar(s5, "Business Model Canvas (BMC)", "04. Business Model")

    table_shape = s5.shapes.add_table(3, 5, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.15))
    table = table_shape.table

    for i in range(5):
        table.columns[i].width = Inches(2.346)

    bmc_headers = ["Key Partners", "Key Activities", "Value Propositions", "Customer Relationships", "Customer Segments"]
    for i, h in enumerate(bmc_headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_FILL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_BLUE
        p.font.name = FONT

    bmc_row1 = [
        "• District Hospitals\n• Private Nursing Homes\n• IMA Doctor Chapters\n• Local Pharmacies",
        "• Clinic Onboarding\n• Real-Time Queue Cloud\n• Patient SMS Dispatch\n• ABDM Compliance",
        "• Zero Waiting Chaos\n• 40% Higher Clinic Flow\n• Live Token on Mobile\n• Zero Hardware Cost",
        "• Self-serve QR Portal\n• Dedicated WhatsApp Support\n• Automated SMS Reminders",
        "• Rural & Town Patients\n• Solo & Polyclinic Doctors\n• Private Nursing Homes\n• District Health Dept"
    ]
    for i, val in enumerate(bmc_row1):
        cell = table.cell(1, i)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.color.rgb = C_MUTED
        p.font.name = FONT

    cell_cost = table.cell(2, 0)
    cell_cost.text = "Cost Structure:\n• Cloud Servers (Next.js/Supabase)\n• SMS/WhatsApp API Gateway\n• Field QR Kits & Operations"
    cell_cost.fill.solid()
    cell_cost.fill.fore_color.rgb = CARD_BG
    p = cell_cost.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_CORAL
    p.font.bold = True
    p.font.name = FONT

    cell_rev = table.cell(2, 2)
    cell_rev.text = "Revenue Streams:\n• Hospital SaaS: ₹999–₹2,999/month\n• Digital Token Convenience: ₹5–₹10\n• Pharmacy & Diagnostic Integrations"
    cell_rev.fill.solid()
    cell_rev.fill.fore_color.rgb = CARD_BG
    p = cell_rev.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_EMERALD
    p.font.bold = True
    p.font.name = FONT

    # ----------------------------------------------------
    # SLIDE 6: MARKET SIZE
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_top_bar(s6, "Market Size (TAM / SAM / SOM)", "05. Market Opportunity")

    mkt_data = [
        ("TAM (India OPD Consultations)", "₹1.8 Lakh Cr", "4.5 Billion annual outpatient consultations across Indian healthcare ecosystem.", C_BLUE),
        ("SAM (Tier-2/3 TS & AP)", "₹850 Crore", "25,000+ small-to-mid private clinics and nursing homes across semi-urban TS & AP.", C_BLUE),
        ("SOM (5 Launch Districts)", "₹4.2 Cr ARR", "350 target clinics across Mahabubabad, Warangal, Khammam, Suryapet, and Karimnagar.", C_EMERALD)
    ]

    for i, (label, val, desc, col) in enumerate(mkt_data):
        x = Inches(0.8 + i * 4.04)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.65), Inches(3.65), Inches(5.15))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = label
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_BLUE

        p = tf.add_paragraph()
        p.text = val
        p.font.name = FONT
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_before = Pt(14)

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(14)

    # ----------------------------------------------------
    # SLIDE 7: COMPETITIVE MATRIX
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_top_bar(s7, "Competitive Advantage Matrix", "06. Market Comparison")

    comp_table_shape = s7.shapes.add_table(6, 4, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.15))
    ct = comp_table_shape.table
    ct.columns[0].width = Inches(3.0)
    ct.columns[1].width = Inches(2.9)
    ct.columns[2].width = Inches(2.9)
    ct.columns[3].width = Inches(2.933)

    comp_headers = ["Key Parameters", "Practo / Apollo 24/7", "Qmatic / Kiosk Hardware", "Skip-Q (Our Solution)"]
    for i, h in enumerate(comp_headers):
        cell = ct.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_FILL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_EMERALD if i == 3 else C_BLUE
        p.font.name = FONT

    comp_rows = [
        ("Live In-Room Token Tracking", "❌ No (Slot-only)", "⚠️ Only on lobby TV", "✅ Real-time on Mobile"),
        ("Hardware Capital Cost", "None", "❌ ₹1.5L+ per kiosk", "✅ ₹0 (Runs on Browser)"),
        ("Tier-2/3 Town Focus", "❌ Metro-Centric", "❌ Corporate Hospitals", "✅ Hyperlocal Bharat"),
        ("Clinic Onboarding Time", "2–3 Weeks", "4+ Weeks", "✅ Under 10 Minutes"),
        ("Pricing for Small Clinics", "❌ High Commission", "❌ Expensive CapEx", "✅ ₹999 / month flat")
    ]

    for r_idx, row in enumerate(comp_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = ct.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = C_EMERALD if c_idx == 3 else (C_WHITE if c_idx == 0 else C_MUTED)
            p.font.bold = (c_idx == 0 or c_idx == 3)
            p.font.name = FONT

    # ----------------------------------------------------
    # SLIDE 8: THE ASK & ROADMAP
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_top_bar(s8, "What We Seek from JNTUH TBI", "07. Incubation Roadmap")

    asks = [
        ("🏛️ 1. Incubation & Clinical Mentorship", "Clinical workflow validation, ABDM integration guidance, and healthcare compliance mentorship from JNTUH TBI ecosystem."),
        ("💰 2. ₹5 Lakhs Pre-Seed Grant", "To fund SMS/WhatsApp gateway scaling, QR onboarding kits, and 50-clinic pilot rollout in Mahabubabad & Warangal."),
        ("🤝 3. Government TVVP Hospital Pilot", "Facilitate trial rollout at Telangana Vaidya Vidhana Parishad (TVVP) Area Hospital in Mahabubabad.")
    ]

    for i, (title, desc) in enumerate(asks):
        y = Inches(1.65 + i * 1.75)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.55))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = C_BLUE
        c.line.width = Pt(1.5)

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.25)
        tf.margin_right = Inches(0.4)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_BLUE

        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(6)

    # ----------------------------------------------------
    # SLIDE 9: THANK YOU & DEMO
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s9)

    c9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    c9.fill.solid()
    c9.fill.fore_color.rgb = CARD_BG
    c9.line.color.rgb = C_EMERALD
    c9.line.width = Pt(2)

    tf = c9.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.6)
    tf.margin_right = Inches(0.6)

    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = FONT
    p.font.size = Pt(52)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = '"Restoring Time and Dignity to Healthcare Waiting Rooms across Bharat"'
    p.font.name = FONT
    p.font.size = Pt(17)
    p.font.italic = True
    p.font.color.rgb = C_WHITE
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Experience the Live Working Prototype Now:\n• 📱 Patient Web App: https://skipq-user.vercel.app\n• 🏢 Hospital Reception Caller: https://skipq-hospital.vercel.app\n• 👑 Super Admin Governance: https://skipq-admin.vercel.app\n• 💻 GitHub Repository: https://github.com/sameir-dev/Skip-Q.git"
    p.font.name = FONT
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_before = Pt(26)

    # Save to both file names
    target1 = "C:/OnlineConsultation/Skip_Q_Pitch_Deck.pptx"
    target2 = "C:/OnlineConsultation/Skip_Q_JNTUH_Presentation_V2.pptx"
    
    prs.save(target1)
    prs.save(target2)
    print(f"Executive Presentation generated successfully at: {target1} and {target2}")

if __name__ == "__main__":
    build_presentation()
